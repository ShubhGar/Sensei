"""
Export API routes
"""
from fastapi import APIRouter, HTTPException, Depends
from fastapi.responses import Response, StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
import json
import io
from datetime import datetime

from app.database import get_db, MeetingSession

router = APIRouter()


@router.get("/pdf/{session_id}")
async def export_pdf(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Export meeting transcript and MOM as PDF"""
    try:
        result = await db.execute(
            select(MeetingSession).where(MeetingSession.id == session_id)
        )
        session = result.scalar_one_or_none()
        
        if not session:
            raise HTTPException(status_code=404, detail="Meeting session not found")
        
        # Import reportlab for PDF generation
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
            from reportlab.lib.enums import TA_CENTER, TA_LEFT
        except ImportError:
            raise HTTPException(status_code=500, detail="PDF generation library not available")
        
        # Create PDF in memory
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter)
        styles = getSampleStyleSheet()
        story = []
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            textColor='#1a1a1a',
            spaceAfter=30,
            alignment=TA_CENTER
        )
        story.append(Paragraph(session.title or "Meeting Transcript", title_style))
        story.append(Spacer(1, 0.2*inch))
        
        # Meeting info
        info_text = f"<b>Session ID:</b> {session.id}<br/>"
        info_text += f"<b>Started:</b> {session.started_at.strftime('%Y-%m-%d %H:%M:%S') if session.started_at else 'N/A'}<br/>"
        if session.ended_at:
            info_text += f"<b>Ended:</b> {session.ended_at.strftime('%Y-%m-%d %H:%M:%S')}<br/>"
        story.append(Paragraph(info_text, styles['Normal']))
        story.append(Spacer(1, 0.3*inch))
        
        # MOM if available
        if session.mom:
            story.append(Paragraph("<b>Minutes of Meeting</b>", styles['Heading2']))
            mom = session.mom if isinstance(session.mom, dict) else json.loads(session.mom) if isinstance(session.mom, str) else {}
            
            if mom.get('summary'):
                story.append(Paragraph("<b>Summary:</b>", styles['Heading3']))
                story.append(Paragraph(mom.get('summary', ''), styles['Normal']))
                story.append(Spacer(1, 0.2*inch))
            
            if mom.get('action_items'):
                story.append(Paragraph("<b>Action Items:</b>", styles['Heading3']))
                for item in mom.get('action_items', []):
                    item_text = f"• {item.get('task', '')}"
                    if item.get('owner'):
                        item_text += f" (Owner: {item.get('owner')})"
                    if item.get('deadline'):
                        item_text += f" (Deadline: {item.get('deadline')})"
                    story.append(Paragraph(item_text, styles['Normal']))
                story.append(Spacer(1, 0.2*inch))
            
            if mom.get('decisions'):
                story.append(Paragraph("<b>Decisions:</b>", styles['Heading3']))
                for decision in mom.get('decisions', []):
                    story.append(Paragraph(f"• {decision.get('decision', '')}", styles['Normal']))
                story.append(Spacer(1, 0.2*inch))
            
            story.append(PageBreak())
        
        # Transcript
        story.append(Paragraph("<b>Full Transcript</b>", styles['Heading2']))
        transcript_text = session.full_transcript or session.transcript or "No transcript available"
        # Split transcript into paragraphs for better formatting
        paragraphs = transcript_text.split('\n\n')
        for para in paragraphs:
            if para.strip():
                story.append(Paragraph(para.strip().replace('\n', '<br/>'), styles['Normal']))
                story.append(Spacer(1, 0.1*inch))
        
        doc.build(story)
        buffer.seek(0)
        
        return Response(
            content=buffer.read(),
            media_type="application/pdf",
            headers={"Content-Disposition": f"attachment; filename=meeting_{session_id}.pdf"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error exporting PDF: {str(e)}")


@router.get("/docx/{session_id}")
async def export_docx(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Export meeting transcript and MOM as DOCX"""
    try:
        result = await db.execute(
            select(MeetingSession).where(MeetingSession.id == session_id)
        )
        session = result.scalar_one_or_none()
        
        if not session:
            raise HTTPException(status_code=404, detail="Meeting session not found")
        
        # Import python-docx for DOCX generation
        try:
            from docx import Document
            from docx.shared import Inches, Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise HTTPException(status_code=500, detail="DOCX generation library not available")
        
        # Create document
        doc = Document()
        
        # Title
        title = doc.add_heading(session.title or "Meeting Transcript", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Meeting info
        doc.add_paragraph(f"Session ID: {session.id}")
        doc.add_paragraph(f"Started: {session.started_at.strftime('%Y-%m-%d %H:%M:%S') if session.started_at else 'N/A'}")
        if session.ended_at:
            doc.add_paragraph(f"Ended: {session.ended_at.strftime('%Y-%m-%d %H:%M:%S')}")
        
        doc.add_paragraph()  # Empty line
        
        # MOM if available
        if session.mom:
            doc.add_heading('Minutes of Meeting', 1)
            mom = session.mom if isinstance(session.mom, dict) else json.loads(session.mom) if isinstance(session.mom, str) else {}
            
            if mom.get('summary'):
                doc.add_heading('Summary', 2)
                doc.add_paragraph(mom.get('summary', ''))
            
            if mom.get('action_items'):
                doc.add_heading('Action Items', 2)
                for item in mom.get('action_items', []):
                    item_text = f"• {item.get('task', '')}"
                    if item.get('owner'):
                        item_text += f" (Owner: {item.get('owner')})"
                    if item.get('deadline'):
                        item_text += f" (Deadline: {item.get('deadline')})"
                    doc.add_paragraph(item_text)
            
            if mom.get('decisions'):
                doc.add_heading('Decisions', 2)
                for decision in mom.get('decisions', []):
                    doc.add_paragraph(f"• {decision.get('decision', '')}")
            
            doc.add_page_break()
        
        # Transcript
        doc.add_heading('Full Transcript', 1)
        transcript_text = session.full_transcript or session.transcript or "No transcript available"
        paragraphs = transcript_text.split('\n\n')
        for para in paragraphs:
            if para.strip():
                doc.add_paragraph(para.strip())
        
        # Save to bytes
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        
        return Response(
            content=buffer.read(),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f"attachment; filename=meeting_{session_id}.docx"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error exporting DOCX: {str(e)}")


@router.get("/json/{session_id}")
async def export_json(
    session_id: str,
    db: AsyncSession = Depends(get_db)
):
    """Export meeting data as JSON"""
    try:
        result = await db.execute(
            select(MeetingSession).where(MeetingSession.id == session_id)
        )
        session = result.scalar_one_or_none()
        
        if not session:
            raise HTTPException(status_code=404, detail="Meeting session not found")
        
        # Build JSON response
        data = {
            "id": session.id,
            "title": session.title,
            "started_at": session.started_at.isoformat() if session.started_at else None,
            "ended_at": session.ended_at.isoformat() if session.ended_at else None,
            "is_active": session.is_active,
            "transcript": session.transcript,
            "full_transcript": session.full_transcript,
            "mom": session.mom if isinstance(session.mom, dict) else json.loads(session.mom) if isinstance(session.mom, str) else session.mom,
            "faqs": session.faqs if isinstance(session.faqs, list) else json.loads(session.faqs) if isinstance(session.faqs, str) else session.faqs,
            "email_draft": session.email_draft if isinstance(session.email_draft, dict) else json.loads(session.email_draft) if isinstance(session.email_draft, str) else session.email_draft,
            "risks": session.risks if isinstance(session.risks, list) else json.loads(session.risks) if isinstance(session.risks, str) else session.risks
        }
        
        json_str = json.dumps(data, indent=2, default=str)
        
        return Response(
            content=json_str,
            media_type="application/json",
            headers={"Content-Disposition": f"attachment; filename=meeting_{session_id}.json"}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error exporting JSON: {str(e)}")
